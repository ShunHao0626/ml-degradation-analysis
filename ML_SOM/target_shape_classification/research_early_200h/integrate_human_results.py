#!/usr/bin/env python3
"""Integrate frozen stage candidates, paper axis correction and one user review.

The review workbook is read only. Image-only opinions remain distinguishable
from numeric 0–200 h stage evidence throughout the resulting delivery.
"""
from __future__ import annotations

import csv
import json
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from early_data import HERE, digest, prepare, read_observations, rows, write_rows
from diagram_stage_classifier import classify, plot, unique_observations, write_point_csv


SOURCE = HERE.parents[1] / "original_curves_2250"
IMAGES = HERE.parents[1] / "data_final"
STAGE = HERE / "diagram_stage_20260925"
REVIEW = HERE / "reviews/outputs/unresolved_only_20260925"
WORKBOOK = REVIEW / "PCE_仅需人工判读15条.xlsx"
OUT = HERE / "final_integrated_20260925"
CLASSES = ("bridge", "hill", "slope", "valley")
HUMAN_CLASSES = {name.title(): name for name in CLASSES}
FORM_STATUS = {"原图可判读", "初期证据仍不足", "不是目标PCE", "需要重新采点", "不确定"}
FORM_CLASS = set(HUMAN_CLASSES) | {"不符合四类", "无法确定"}


def locked_review():
    before = digest(WORKBOOK)
    expected = rows(REVIEW / "manual_queue.csv")
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = book["待判读15条"]
    opinions = []
    for position, expected_row in enumerate(expected, 7):
        fid = sheet.cell(position, 2).value
        if fid != expected_row["curve_id"]:
            raise ValueError(f"review ID mismatch at row {position}: {fid}")
        status = sheet.cell(position, 6).value
        label = sheet.cell(position, 7).value
        note = sheet.cell(position, 8).value or ""
        if status not in FORM_STATUS or label not in FORM_CLASS:
            raise ValueError(f"invalid or incomplete opinion: {fid}: {status}/{label}")
        if status != "原图可判读" and label in HUMAN_CLASSES:
            raise ValueError("four-class label without readable source evidence: " + fid)
        opinions.append(dict(curve_id=fid, review_row=position,
                             original_image_readability=status,
                             human_image_label=label, human_note=note,
                             numeric_strict_points=expected_row["strict_unique_times"],
                             target_series=expected_row["target_series"],
                             workbook_sha256=before))
    book.close()
    if digest(WORKBOOK) != before:
        raise ValueError("review workbook changed during import")
    if len(opinions) != 15 or len({r["curve_id"] for r in opinions}) != 15:
        raise ValueError("review count changed")
    return opinions, before


def write_raw_points(path, source_row):
    points = read_observations(source_row, SOURCE)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["hour", "raw_pce", "source_row", "region", "normalized_pce"])
        for hour, value, source_index in points:
            if hour >= 0:
                writer.writerow([f"{hour:.12g}", f"{value:.12g}", source_index,
                                 "0_200h" if hour <= 200 else "after_200h", ""])


def copy_existing(stage_row, record):
    for field, suffix in (("curve_csv_path", ".csv"), ("plot_path", ".svg")):
        source = STAGE / stage_row[field]
        target = OUT / "classes" / record["integrated_class"] / (record["curve_id"] + suffix)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        record["curve_csv" if suffix == ".csv" else "curve_image"] = str(target.relative_to(OUT))


def recover_paper_pce(source, record, expected):
    revised = dict(source, y_kind="pce")
    item, error = prepare(revised, SOURCE,
                          dict(window_h=200, boundary_limit_h=250, boundary_max_points=3))
    if error:
        raise ValueError(record["curve_id"] + ": " + error)
    t, y = unique_observations(item["points"], item["divisor"])
    stage = classify(t, y)
    if stage["candidate_class"] != expected["automatic_class"]:
        raise ValueError("paper-based class changed: " + record["curve_id"])
    record.update(integrated_class=stage["candidate_class"],
                  evidence_status=stage["evidence_status"],
                  classification_origin="paper_axis_check_plus_automatic_stage_rule",
                  numeric_strict_points=item["unique_window_times"],
                  normalization_divisor=item["divisor"],
                  primary_turn_h=stage["primary_turn_h"],
                  secondary_turn_h=stage["secondary_turn_h"],
                  decay_transition_h=stage["decay_transition_h"],
                  initial_amplitude=stage["initial_amplitude"],
                  reversal_amplitude=stage["reversal_amplitude"],
                  initial_rate_per_h=stage["initial_rate_per_h"],
                  reversal_rate_per_h=stage["reversal_rate_per_h"],
                  slow_decay_rate_per_h=stage["slow_decay_rate_per_h"],
                  post_200_pattern=stage["post_200_pattern"],
                  post_200_matches_diagram=stage["post_200_matches_diagram"],
                  decision_detail=expected["evidence"])
    folder = OUT / "classes" / record["integrated_class"]
    csv_path = folder / (record["curve_id"] + ".csv")
    image_path = folder / (record["curve_id"] + ".svg")
    write_point_csv(csv_path, item["points"], item["divisor"])
    plot(image_path, t, y, dict(stage, curve_id=record["curve_id"]))
    record["curve_csv"] = str(csv_path.relative_to(OUT))
    record["curve_image"] = str(image_path.relative_to(OUT))


def write_image_review_files(source, record):
    fid = record["curve_id"]
    folder = (OUT / "classes" / record["integrated_class"] / "image_only"
              if record["integrated_class"] else OUT / "nonconforming")
    folder.mkdir(parents=True, exist_ok=True)
    raw = folder / (fid + "_raw.csv")
    original = folder / (fid + "_source" + Path(source["source_image"]).suffix.lower())
    write_raw_points(raw, source)
    shutil.copy2(IMAGES / source["source_image"], original)
    record["curve_csv"] = str(raw.relative_to(OUT))
    record["curve_image"] = str(original.relative_to(OUT))


def write_gallery(records):
    fields = ("curve_id", "source_group", "integrated_class", "evidence_status",
              "classification_origin", "human_image_label", "target_series",
              "numeric_strict_points", "curve_csv", "curve_image", "source_image")
    payload = [{key: r.get(key, "") for key in fields} for r in records]
    data = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    page = r'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>PCE 初期四类整合结果</title>
<style>body{font:14px system-ui;margin:22px;color:#263840}main{display:flex;gap:20px}aside{width:42%;height:82vh;overflow:auto}article{flex:1;position:sticky;top:15px;height:82vh}input,select{padding:7px;margin-right:8px}table{border-collapse:collapse;width:100%}td,th{padding:6px;border-bottom:1px solid #ddd;text-align:left}tr:hover{background:#eaf2f4;cursor:pointer}img{width:100%;max-height:500px;object-fit:contain;border:1px solid #ddd}a{margin-right:12px}</style>
<h1>PCE 初期四类整合结果</h1><p>四类包含自动候选和 9 条原图人工意见。原图人工意见在 0–200 h 只有 0 或 1 个数字时间点，不代表数值阶段已证实。非目标物理量和非四类单独保留。</p><p><a href="result/index.html">四类图文总览</a> · <a href="../evaluation/strict_numeric_coverage_20260925/README.md">严格 0–200 h 数值证据范围</a> · <a href="../evaluation/stage_sensitivity_20260925/README.md">规则稳健性审计</a></p>
<input id="search" placeholder="曲线 ID、DOI 或系列"><select id="cls"><option value="">全部类别</option><option>bridge</option><option>hill</option><option>slope</option><option>valley</option><option value="none">无四类</option></select><select id="origin"><option value="">全部来源</option><option value="single_user_source_image">原图人工判读</option><option value="automatic_stage_rule">自动规则</option><option value="paper_axis_check_plus_automatic_stage_rule">论文核轴后自动规则</option></select><span id="count"></span>
<main><aside><table><thead><tr><th>ID</th><th>类别</th><th>证据状态</th><th>来源</th></tr></thead><tbody id="body"></tbody></table></aside><article><h2 id="title">选择曲线</h2><img id="chart"><div id="links"></div><p id="info"></p></article></main>
<script>const rows=__DATA__,body=document.querySelector('#body');function show(r){document.querySelector('#title').textContent=r.curve_id+' | '+(r.integrated_class||'无四类');document.querySelector('#chart').src=r.curve_image||('../../../data_final/'+r.source_image);document.querySelector('#info').textContent='来源：'+r.classification_origin+'；0–200 h 数字时间点：'+(r.numeric_strict_points||'—')+'；人工意见：'+(r.human_image_label||'—')+'；目标系列：'+(r.target_series||'—');let links=document.querySelector('#links');links.replaceChildren();for(const [label,url] of [['曲线数据',r.curve_csv],['本次曲线图',r.curve_image],['论文原图',r.source_image?'../../../data_final/'+r.source_image:'']])if(url){let a=document.createElement('a');a.href=url;a.textContent=label;a.target='_blank';links.append(a)}}function render(){body.replaceChildren();let q=document.querySelector('#search').value.toLowerCase(),c=document.querySelector('#cls').value,o=document.querySelector('#origin').value;let shown=rows.filter(r=>(!c||(c==='none'?!r.integrated_class:r.integrated_class===c))&&(!o||r.classification_origin===o)&&(!q||(r.curve_id+r.source_group+r.target_series).toLowerCase().includes(q)));document.querySelector('#count').textContent=shown.length+' / '+rows.length;for(const r of shown){let tr=document.createElement('tr');for(const key of ['curve_id','integrated_class','evidence_status','classification_origin']){let td=document.createElement('td');td.textContent=r[key];tr.append(td)}tr.onclick=()=>show(r);body.append(tr)}if(shown.length)show(shown[0])}for(const id of ['search','cls','origin'])document.querySelector('#'+id).addEventListener('input',render);render();</script>'''.replace("__DATA__", data)
    (OUT / "all_curves_index.html").write_text(page, encoding="utf-8")


def main():
    opinions, workbook_sha = locked_review()
    stage = {r["curve_id"]: r for r in rows(STAGE / "stage_features.csv")}
    original_delivery = {r["curve_id"]: r for r in rows(HERE / "delivery/class_index.csv")}
    triage = {r["curve_id"]: r for r in rows(REVIEW / "triage_audit.csv")}
    frozen = rows(HERE / "manifests/frozen_inputs.csv")
    canonical_ids = sorted({r["canonical_id"] for r in frozen})
    by_file = {r["file_id"]: r for r in frozen}
    aliases = defaultdict(list)
    for r in frozen:
        aliases[r["canonical_id"]].append(r["file_id"])
    human = {r["curve_id"]: r for r in opinions}
    if len(stage) != 2205 or len(triage) != 41 or len(canonical_ids) != 2246:
        raise ValueError("source coverage changed")
    if set(stage) & set(triage) or set(stage) | set(triage) != set(canonical_ids):
        raise ValueError("stage and triage IDs do not partition canonical curves")
    if set(human) != {fid for fid, r in triage.items() if r["disposition"] == "manual_source_image_check"}:
        raise ValueError("human review IDs differ from manual queue")
    OUT.mkdir(parents=True, exist_ok=True)
    write_rows(OUT / "human_opinions_locked.csv", opinions, list(opinions[0]))
    records = []
    for fid in canonical_ids:
        source = by_file[fid]
        record = dict(curve_id=fid, file_ids=";".join(aliases[fid]),
                      source_group=source["source_group"], split=source["split"],
                      source_csv=source["source_csv"], source_sha256=source["source_sha256"],
                      source_image=source["source_image"], integrated_class="",
                      evidence_status="", classification_origin="", human_image_label="",
                      human_review_status="", human_note="", target_series="",
                      numeric_strict_points="", normalization_divisor="",
                      primary_turn_h="", secondary_turn_h="", decay_transition_h="",
                      initial_amplitude="", reversal_amplitude="", initial_rate_per_h="",
                      reversal_rate_per_h="", slow_decay_rate_per_h="",
                      post_200_pattern="", post_200_matches_diagram="", decision_detail="",
                      curve_csv="", curve_image="")
        if fid in stage:
            previous = stage[fid]
            record.update(integrated_class=previous["candidate_class"],
                          evidence_status=previous["evidence_status"],
                          classification_origin="automatic_stage_rule",
                          numeric_strict_points=previous["strict_points"],
                          normalization_divisor=original_delivery[fid]["normalization_divisor"],
                          primary_turn_h=previous["primary_turn_h"],
                          secondary_turn_h=previous["secondary_turn_h"],
                          decay_transition_h=previous["decay_transition_h"],
                          initial_amplitude=previous["initial_amplitude"],
                          reversal_amplitude=previous["reversal_amplitude"],
                          initial_rate_per_h=previous["initial_rate_per_h"],
                          reversal_rate_per_h=previous["reversal_rate_per_h"],
                          slow_decay_rate_per_h=previous["slow_decay_rate_per_h"],
                          post_200_pattern=previous["post_200_pattern"],
                          post_200_matches_diagram=previous["post_200_matches_diagram"],
                          decision_detail=previous["reason"])
            copy_existing(previous, record)
        else:
            decision = triage[fid]
            record["target_series"] = decision["target_series"]
            if decision["disposition"] == "auto_resolved_from_paper_and_csv":
                recover_paper_pce(source, record, decision)
            elif decision["disposition"] == "manual_source_image_check":
                opinion = human[fid]
                record.update(human_image_label=opinion["human_image_label"],
                              human_review_status=opinion["original_image_readability"],
                              human_note=opinion["human_note"],
                              numeric_strict_points=decision["strict_unique_times"],
                              classification_origin="single_user_source_image",
                              decision_detail="user reviewed original figure; numeric 0–200 h stage has fewer than two times")
                if opinion["human_image_label"] in HUMAN_CLASSES:
                    record.update(integrated_class=HUMAN_CLASSES[opinion["human_image_label"]],
                                  evidence_status="human_image_only_numeric_stage_insufficient")
                elif opinion["human_image_label"] == "不符合四类":
                    record["evidence_status"] = "human_nonconforming_image_only"
                else:
                    record["evidence_status"] = "human_undetermined_image_only"
                write_image_review_files(source, record)
            else:
                record.update(evidence_status="out_of_scope", classification_origin="source_axis_audit",
                              decision_detail=decision["evidence"])
        records.append(record)
    fields = list(records[0])
    write_rows(OUT / "canonical_index.csv", records, fields)
    write_rows(OUT / "class_index.csv", [r for r in records if r["integrated_class"]], fields)
    write_rows(OUT / "nonconforming.csv", [r for r in records if r["human_image_label"] == "不符合四类"], fields)
    write_rows(OUT / "no_numeric_strict_support.csv", [r for r in records if r["classification_origin"] == "single_user_source_image"], fields)
    write_rows(OUT / "out_of_scope.csv", [r for r in records if r["evidence_status"] == "out_of_scope"], fields)
    by_id = {r["curve_id"]: r for r in records}
    all_files = []
    for row in frozen:
        assigned = by_id[row["canonical_id"]]
        all_files.append(dict(file_id=row["file_id"], canonical_id=row["canonical_id"],
                              source_csv=row["source_csv"], source_sha256=row["source_sha256"],
                              source_group=row["source_group"],
                              integrated_class=assigned["integrated_class"],
                              evidence_status=assigned["evidence_status"],
                              classification_origin=assigned["classification_origin"],
                              curve_csv=assigned["curve_csv"], curve_image=assigned["curve_image"]))
    write_rows(OUT / "all_files.csv", all_files, list(all_files[0]))
    write_gallery(records)
    counts = Counter(r["integrated_class"] for r in records if r["integrated_class"])
    origins = Counter(r["classification_origin"] for r in records)
    expected = dict(bridge=88, hill=180, slope=1872, valley=76)
    if dict(counts) != expected or len(all_files) != 2250:
        raise ValueError(f"unexpected final coverage: {counts}, {len(all_files)}")
    for record in records:
        if record["integrated_class"]:
            for field in ("curve_csv", "curve_image"):
                if not (OUT / record[field]).is_file():
                    raise ValueError("missing classified artifact: " + record["curve_id"] + " " + field)
    summary = dict(protocol="integrated_200h_20260925", original_source_root=str(SOURCE),
                   human_workbook_sha256=workbook_sha,
                   canonical_curves=len(records), original_files=len(all_files),
                   classes=dict(counts), origins=dict(origins),
                   human_four_class=sum(r["classification_origin"] == "single_user_source_image" and bool(r["integrated_class"]) for r in records),
                   human_nonconforming=sum(r["human_image_label"] == "不符合四类" for r in records),
                   out_of_scope=sum(r["evidence_status"] == "out_of_scope" for r in records),
                   numeric_0_200h_stage_absent_for_human=15,
                   claim="Integrated candidates with provenance; no independent accuracy estimate")
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
